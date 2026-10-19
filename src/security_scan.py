"""security_check — ports & server security (doc/SECURITY_TOOLS_DECISION.md §6, options D + A).

Close to the machine, no model involved — every fact is measured:
  D  from outside: raw TCP connects to ~40 well-known ports, the banner each open service sends (web
     ports: one `HEAD /` request), and Python's own TLS check of HTTPS certificates
  A  this PC (the MulteAgent server): the OS connection table (psutil) — listening ports, owning
     program, reachable from the network or local only — plus Windows Firewall profiles, Remote
     Desktop (registry) and SMBv1 (fixed PowerShell commands)
Every finding comes from the fixed rule table below, with the reason and the fix as advice.

Same safety as network_scan: private ranges only, NETWORK_SCAN_MAX_HOSTS, a time limit, Stop,
SECURITY_SCAN_ENABLED=false switches it off; the worker lets only admins run it. Read-only: no
exploits, no password guessing, nothing changed; each open port is read once.
"""

from __future__ import annotations

import ipaddress
import json
import logging
import re
import socket
import ssl
import subprocess
import sys
import time
from collections.abc import Callable
from concurrent.futures import ThreadPoolExecutor, as_completed
from dataclasses import dataclass, field

from src.config import Config
from src.net_scan import PRIVATE_NETS, ScanError, local_subnet, validate_subnet

logger = logging.getLogger(__name__)

TIME_LIMIT_SECONDS = 120
WORKERS = 64
TCP_TIMEOUT = 0.5
BANNER_TIMEOUT = 2.0
MANY_PORTS = 12
MAX_SERVICE_ROWS = 30          # longer service lists go to the Excel file only
_NO_WINDOW = getattr(subprocess, "CREATE_NO_WINDOW", 0)

HIGH, MEDIUM, LOW, INFO = "high", "medium", "low", "info"
RISK_ORDER = {HIGH: 0, MEDIUM: 1, LOW: 2, INFO: 3}
RISK_LABEL = {HIGH: "🔴 High", MEDIUM: "🟠 Medium", LOW: "🟡 Low", INFO: "⚪ Info"}


@dataclass(frozen=True)
class PortRule:
    name: str
    risk: str
    why: str = ""
    fix: str = ""


# the ports checked from outside, and what an open one means (risk for a port reachable on the LAN)
PORTS: dict[int, PortRule] = {
    21: PortRule("FTP", HIGH, "FTP sends passwords and files unencrypted", "Use SFTP (over SSH) or FTPS; stop the FTP service if unused"),
    22: PortRule("SSH", INFO),
    23: PortRule("Telnet", HIGH, "Telnet sends passwords in clear text", "Switch Telnet off and use SSH"),
    25: PortRule("SMTP", INFO),
    53: PortRule("DNS", INFO),
    80: PortRule("HTTP", INFO),
    110: PortRule("POP3", MEDIUM, "Mail login without encryption", "Use POP3S (995) and close 110"),
    135: PortRule("MS-RPC", MEDIUM, "Windows RPC reachable — often used to move between PCs in an attack", "Allow RPC only from admin PCs in Windows Firewall"),
    139: PortRule("NetBIOS", MEDIUM, "Old Windows name/file service reachable", "Disable NetBIOS over TCP/IP if no old devices need it"),
    143: PortRule("IMAP", MEDIUM, "Mail login without encryption", "Use IMAPS (993) and close 143"),
    389: PortRule("LDAP", LOW, "Directory queries without encryption", "Prefer LDAPS (636) or LDAP signing"),
    443: PortRule("HTTPS", INFO),
    445: PortRule("SMB", MEDIUM, "Windows file sharing reachable — the path ransomware spreads by", "Share only what is needed; block 445 for devices that need no shares"),
    465: PortRule("SMTPS", INFO),
    554: PortRule("RTSP (camera)", MEDIUM, "Camera video stream reachable", "Set a strong camera password; keep cameras in their own network (VLAN)"),
    587: PortRule("SMTP submission", INFO),
    636: PortRule("LDAPS", INFO),
    993: PortRule("IMAPS", INFO),
    995: PortRule("POP3S", INFO),
    1433: PortRule("SQL Server", HIGH, "Database reachable from the network", "Allow only the app servers that need it (firewall), or bind to localhost"),
    1521: PortRule("Oracle DB", HIGH, "Database reachable from the network", "Allow only the app servers that need it (firewall)"),
    1883: PortRule("MQTT", MEDIUM, "IoT messages without encryption, often without a password", "Require a login and use MQTT over TLS (8883)"),
    2049: PortRule("NFS", MEDIUM, "Network file share reachable", "Export only to the hosts that need it"),
    2375: PortRule("Docker API", HIGH, "Unencrypted Docker API = full control of the host", "Close 2375; use the local socket or TLS (2376)"),
    3000: PortRule("Web app (3000)", LOW, "Development web server reachable", "Run it behind a reverse proxy or bind to localhost"),
    3306: PortRule("MySQL", HIGH, "Database reachable from the network", "Bind to localhost or allow only the app servers (firewall)"),
    3389: PortRule("RDP", HIGH, "Remote Desktop reachable — the #1 target of password-guessing attacks", "Allow RDP only via VPN or from admin PCs; require Network Level Authentication"),
    5000: PortRule("Web app (5000)", LOW, "Web service reachable", "Check it is meant to be shared"),
    5432: PortRule("PostgreSQL", HIGH, "Database reachable from the network", "Bind to localhost or allow only the app servers (pg_hba.conf + firewall)"),
    5601: PortRule("Kibana", MEDIUM, "Log dashboard reachable", "Put it behind a login / reverse proxy"),
    5900: PortRule("VNC", MEDIUM, "Remote screen reachable; often weak passwords", "Use VNC only through SSH or VPN"),
    5985: PortRule("WinRM (HTTP)", MEDIUM, "Remote management without encryption", "Use WinRM over HTTPS (5986) and only from admin PCs"),
    5986: PortRule("WinRM (HTTPS)", LOW, "Remote management reachable", "Allow only from admin PCs"),
    6379: PortRule("Redis", HIGH, "Redis is often reachable without a password", "Bind to localhost and set requirepass"),
    8000: PortRule("HTTP-alt (8000)", LOW, "Extra web service", "Check it is meant to be shared"),
    8080: PortRule("HTTP-alt (8080)", LOW, "Extra web service", "Check it is meant to be shared"),
    8443: PortRule("HTTPS-alt (8443)", INFO),
    8888: PortRule("HTTP-alt (8888)", LOW, "Extra web service", "Check it is meant to be shared"),
    9100: PortRule("Printer (JetDirect)", LOW, "Anyone on the network can send print jobs", "Restrict printing to the print server"),
    9200: PortRule("Elasticsearch", HIGH, "Search database reachable, often without a login", "Bind to localhost or enable security (login + TLS)"),
    10000: PortRule("Webmin", MEDIUM, "Server admin panel reachable", "Allow only from admin PCs"),
    11211: PortRule("Memcached", HIGH, "Cache reachable without a login", "Bind to localhost"),
    27017: PortRule("MongoDB", HIGH, "Database reachable, sometimes without a login", "Bind to localhost and enable authentication"),
}
HTTP_PORTS = {80, 3000, 5000, 5601, 8000, 8080, 8888, 9200, 10000}
TLS_PORTS = {443, 465, 636, 993, 995, 5986, 8443}
TLS_HTTP_PORTS = {443, 8443}
THIS_PC_RE = re.compile(
    r"\b(this|my|mera|meri|apna|apne)\s+(pc|computer|machine|laptop|server|system)\b|\blocalhost\b|\b127\.0\.0\.1\b", re.I)
_IP_RE = re.compile(r"\b(\d{1,3}(?:\.\d{1,3}){3})(?:\s*/\s*(\d{1,2}))?\b")


@dataclass
class Finding:
    risk: str
    host: str
    port: int | None
    title: str
    why: str
    fix: str


@dataclass
class Service:
    port: int
    name: str
    product: str = ""          # "OpenSSH 8.9p1", "nginx/1.18.0"
    tls: str = ""              # certificate status for TLS ports


@dataclass
class Listener:
    address: str
    port: int
    program: str
    network: bool              # reachable from the network (not only 127.0.0.1 / ::1)


@dataclass
class LocalView:
    listeners: list[Listener] = field(default_factory=list)
    firewall: dict[str, bool] = field(default_factory=dict)       # profile → enabled
    rdp_on: bool | None = None
    smb1_on: bool | None = None
    notes: list[str] = field(default_factory=list)
    api_port: int | None = None                                   # to name MulteAgent's own port


@dataclass
class HostReport:
    ip: str
    name: str = ""
    vendor: str = ""
    services: list[Service] = field(default_factory=list)
    local: LocalView | None = None          # set for this PC (checked from the inside)


@dataclass
class SecurityResult:
    target: str
    hosts: list[HostReport]
    findings: list[Finding]
    seconds: float
    stopped: bool = False
    timed_out: bool = False


# ── what to check ────────────────────────────────────────────────────────
def own_ips() -> set[str]:
    import psutil

    return {a.address for addrs in psutil.net_if_addrs().values() for a in addrs if a.family == socket.AF_INET}


def resolve_target(target: str | None, query: str, config: Config) -> tuple[str, list[str] | None, bool]:
    """→ (label, hosts, this_pc). hosts=None means: find the devices of the network in `label` first."""
    text = (target or "").strip()
    if THIS_PC_RE.search(text) or (not text and THIS_PC_RE.search(query or "")):
        return "this PC", [], True
    m = _IP_RE.search(text) or _IP_RE.search(query or "")
    if m and not m.group(2):                                    # one address
        try:
            ip = ipaddress.ip_address(m.group(1))
        except ValueError as exc:
            raise ScanError(f"not an IP address: {m.group(1)}") from exc
        if not any(ip in net for net in PRIVATE_NETS):
            raise ScanError(f"{ip} is not a private address — only 10.x, 172.16–31.x and 192.168.x can be checked")
        if str(ip) in own_ips():
            return "this PC", [], True
        return str(ip), [str(ip)], False
    subnet = f"{m.group(1)}/{m.group(2)}" if m else local_subnet()
    net = validate_subnet(subnet, config.network_scan_max_hosts)
    return str(net), None, False


# ── D: from outside ──────────────────────────────────────────────────────
def _port_open(ip: str, port: int) -> bool:
    try:
        with socket.create_connection((ip, port), timeout=TCP_TIMEOUT):
            return True
    except OSError:
        return False


def _clean(data: bytes) -> str:
    text = data.decode("latin-1", "replace")
    return re.sub(r"[^\x20-\x7e]+", " ", text).strip()[:160]


def _head_request(ip: str) -> bytes:
    return f"HEAD / HTTP/1.0\r\nHost: {ip}\r\nUser-Agent: MulteAgent-security-check\r\n\r\n".encode()


def _read_banner(ip: str, port: int) -> str:
    """What the service says first, or (web ports) the answer to one HEAD request."""
    try:
        with socket.create_connection((ip, port), timeout=BANNER_TIMEOUT) as sock:
            sock.settimeout(BANNER_TIMEOUT)
            if port in HTTP_PORTS:
                sock.sendall(_head_request(ip))
            data = sock.recv(512)
    except OSError:
        return ""
    return data.decode("latin-1", "replace") if port in HTTP_PORTS else _clean(data)


def _tls_http_banner(ip: str, port: int) -> str:
    ctx = ssl.create_default_context()
    ctx.check_hostname, ctx.verify_mode = False, ssl.CERT_NONE      # only to read the Server header
    try:
        with socket.create_connection((ip, port), timeout=BANNER_TIMEOUT) as raw, \
                ctx.wrap_socket(raw, server_hostname=None) as tls:
            tls.sendall(_head_request(ip))
            return tls.recv(512).decode("latin-1", "replace")
    except (ssl.SSLError, OSError):
        return ""


def product_from_banner(port: int, banner: str) -> str:
    """'SSH-2.0-OpenSSH_8.9p1 Ubuntu-3' → 'OpenSSH 8.9p1'; an HTTP answer → its Server header."""
    if not banner:
        return ""
    if m := re.search(r"SSH-[\d.]+-([\w.\-]+)", banner):
        software = m.group(1).replace("_", " ")
        # some devices announce a random name to hide their software (e.g. 'o4o3e') — say so, it is no version
        if not re.search(r"\d+\.\d+|openssh|dropbear|rosssh|cisco|libssh", software, re.I):
            return f"unknown software (announces '{software}')"
        return software
    if m := re.search(r"(?im)^server:\s*(.+?)\s*$", banner):
        return m.group(1)[:80]
    if m := re.match(r"RFB (\d{3}\.\d{3})", banner):
        return f"VNC protocol {m.group(1)}"
    if port == 3306 and (m := re.search(r"(\d+\.\d+\.\d+[\w.\-]*)", banner)):
        return f"MySQL {m.group(1)}"
    if m := re.match(r"(?:220|\+OK|\* OK)[ -](.+)", banner):     # FTP / SMTP / POP3 / IMAP greetings
        return m.group(1).strip()[:80]
    if m := re.match(r"HTTP/[\d.]+ (\d{3})", banner):
        return f"web server (answer {m.group(1)}, no Server header)"
    return ""


def classify_tls_error(message: str) -> tuple[str, str | None]:
    low = message.lower()
    if "expired" in low:
        return "certificate EXPIRED", MEDIUM
    if "self-signed" in low or "self signed" in low:
        return "self-signed certificate (browsers warn; fine for internal tools)", LOW
    return f"certificate not trusted ({message[:60]})", LOW


def tls_status(ip: str, port: int) -> tuple[str, str | None]:
    """→ (text, risk or None). Expired / self-signed / untrusted come from Python's own TLS check."""
    ctx = ssl.create_default_context()
    ctx.check_hostname = False                 # we connect by IP; the name check would always fail
    try:
        with socket.create_connection((ip, port), timeout=BANNER_TIMEOUT) as raw, \
                ctx.wrap_socket(raw, server_hostname=None) as tls:
            cert = tls.getpeercert() or {}
    except ssl.SSLCertVerificationError as exc:
        return classify_tls_error(exc.verify_message or str(exc))
    except (ssl.SSLError, OSError) as exc:
        return f"TLS handshake failed ({type(exc).__name__})", None
    if not cert.get("notAfter"):
        return "certificate valid", None
    expires = ssl.cert_time_to_seconds(cert["notAfter"])
    days = int((expires - time.time()) // 86400)
    until = time.strftime("%Y-%m-%d", time.gmtime(expires))
    if days < 30:
        return f"certificate expires in {days} days ({until})", LOW
    return f"certificate valid until {until}", None


def check_service(ip: str, port: int) -> Service:
    rule = PORTS.get(port, PortRule(f"port {port}", INFO))
    service = Service(port=port, name=rule.name)
    if port in TLS_PORTS:
        service.tls, _ = tls_status(ip, port)
        if port in TLS_HTTP_PORTS:
            service.product = product_from_banner(port, _tls_http_banner(ip, port))
    else:
        service.product = product_from_banner(port, _read_banner(ip, port))
    return service


# ── A: this PC ───────────────────────────────────────────────────────────
def _powershell(command: str) -> str:
    """One fixed, read-only PowerShell command (never built from user text)."""
    out = subprocess.run(["powershell", "-NoProfile", "-NonInteractive", "-Command", command],
                         capture_output=True, text=True, timeout=20, creationflags=_NO_WINDOW)
    if out.returncode != 0:
        raise RuntimeError(((out.stderr or "failed").strip().splitlines() or ["failed"])[0][:120])
    return out.stdout


def parse_firewall(json_text: str) -> dict[str, bool]:
    data = json.loads(json_text or "[]")
    items = data if isinstance(data, list) else [data]
    # Enabled comes back as 1/0, true/false or "True" depending on the PowerShell version
    return {str(p["Name"]): str(p.get("Enabled")).lower() in {"1", "true"} for p in items if p.get("Name")}


def local_view() -> LocalView:
    import psutil

    view = LocalView()
    # one row per (port, program): Windows lists IPv4 (0.0.0.0) and IPv6 (::) separately, and NetBIOS
    # once per network adapter — merged, so each port is one finding
    merged: dict[tuple[int, str], list[str]] = {}
    for conn in psutil.net_connections(kind="inet"):
        if conn.status != psutil.CONN_LISTEN or not conn.laddr:
            continue
        try:
            program = psutil.Process(conn.pid).name() if conn.pid else "system"
        except (psutil.Error, OSError):
            program = "system"
        addresses = merged.setdefault((conn.laddr.port, program), [])
        if conn.laddr.ip not in addresses:
            addresses.append(conn.laddr.ip)
    for (port, program), addresses in merged.items():
        network = any(a not in {"127.0.0.1", "::1"} for a in addresses)
        view.listeners.append(Listener(", ".join(sorted(addresses)), port, program, network))
    view.listeners.sort(key=lambda listener: (not listener.network, listener.port))
    if sys.platform != "win32":
        view.notes.append("Firewall, Remote Desktop and SMB checks are for Windows; this server is not Windows.")
        return view
    try:
        view.firewall = parse_firewall(_powershell("Get-NetFirewallProfile | Select-Object Name,Enabled | ConvertTo-Json"))
    except Exception as exc:  # noqa: BLE001
        view.notes.append(f"Windows Firewall could not be read: {exc}")
    try:
        import winreg

        with winreg.OpenKey(winreg.HKEY_LOCAL_MACHINE, r"System\CurrentControlSet\Control\Terminal Server") as key:
            view.rdp_on = winreg.QueryValueEx(key, "fDenyTSConnections")[0] == 0
    except OSError as exc:
        view.notes.append(f"Remote Desktop setting could not be read: {exc}")
    try:
        view.smb1_on = _powershell("(Get-SmbServerConfiguration).EnableSMB1Protocol").strip().lower() == "true"
    except Exception as exc:  # noqa: BLE001
        view.notes.append(f"SMBv1 setting could not be read (may need administrator): {exc}")
    return view


# ── the rules ────────────────────────────────────────────────────────────
def host_findings(host: HostReport) -> list[Finding]:
    """Findings from the fixed rules for one host seen from outside."""
    out: list[Finding] = []
    ports = {s.port for s in host.services}
    for s in host.services:
        rule = PORTS.get(s.port)
        if rule and rule.risk != INFO:
            out.append(Finding(rule.risk, host.ip, s.port, f"{rule.name} open", rule.why, rule.fix))
        if "EXPIRED" in s.tls:
            out.append(Finding(MEDIUM, host.ip, s.port, "Certificate expired", "Browsers and apps refuse or warn",
                               "Renew the certificate"))
        elif any(word in s.tls for word in ("expires in", "self-signed", "not trusted")):
            out.append(Finding(LOW, host.ip, s.port, s.tls.split(" (")[0].capitalize(),
                               "Users learn to click through certificate warnings", "Use a certificate from a trusted CA"))
        if m := re.search(r"OpenSSH[ _](\d+)\.(\d+)", s.product):
            if (int(m.group(1)), int(m.group(2))) < (8, 0):
                out.append(Finding(LOW, host.ip, s.port, f"Old SSH server ({s.product})",
                                   "OpenSSH older than 8.0 (2019) misses years of security fixes", "Update the OS packages"))
    if 80 in ports and not ports & {443, 8443}:
        out.append(Finding(MEDIUM, host.ip, 80, "Web page without HTTPS",
                           "Logins on this page travel unencrypted", "Enable HTTPS (a certificate), redirect 80 → 443"))
    if len(ports) > MANY_PORTS:
        out.append(Finding(LOW, host.ip, None, f"{len(ports)} services open",
                           "Every open service is something to patch and protect", "Switch off what is not needed"))
    return out


def local_findings(view: LocalView, label: str = "this PC") -> list[Finding]:
    out: list[Finding] = []
    for profile, enabled in view.firewall.items():
        if not enabled:
            out.append(Finding(HIGH, label, None, f"Windows Firewall OFF ({profile} profile)",
                               "Every listening service is reachable", "Turn the firewall on: Windows Security → Firewall"))
    if view.rdp_on:
        out.append(Finding(MEDIUM, label, 3389, "Remote Desktop is on", "A target for password guessing",
                           "Turn it off if unused; else allow only from admin PCs / VPN, with Network Level Authentication"))
    if view.smb1_on:
        out.append(Finding(HIGH, label, 445, "SMBv1 is on", "The protocol WannaCry spread through",
                           "Disable-WindowsOptionalFeature -Online -FeatureName SMB1Protocol"))
    for listener in view.listeners:
        rule = PORTS.get(listener.port)
        if listener.network and rule and rule.risk in (HIGH, MEDIUM):
            out.append(Finding(rule.risk, label, listener.port,
                               f"{rule.name} listening on the network ({listener.program})",
                               f"{rule.why}; Windows Firewall decides whether other devices get through", rule.fix))
    return out


def _order(f: Finding) -> tuple:
    return (RISK_ORDER[f.risk], f.host, f.port or 0)


def _host_key(report: HostReport) -> tuple:
    try:
        return (0, ipaddress.ip_address(report.ip))
    except ValueError:
        return (1, report.ip)


# ── the check ────────────────────────────────────────────────────────────
def check(
    target: str | None,
    query: str,
    config: Config,
    *,
    should_continue: Callable[[], bool] = lambda: True,
    progress: Callable[[str], None] = lambda message: None,
) -> SecurityResult:
    if not config.security_scan_enabled:
        raise ScanError("the security check is switched off (SECURITY_SCAN_ENABLED=false)")
    started = time.monotonic()
    label, hosts, this_pc = resolve_target(target, query, config)
    stopped = timed_out = False

    if this_pc:
        progress("Reading this PC's ports, firewall and settings…")
        view = local_view()
        view.api_port = config.api_port
        report = HostReport(ip="this PC", name=socket.gethostname(), local=view)
        return SecurityResult(label, [report], sorted(local_findings(view), key=_order),
                              round(time.monotonic() - started, 1))

    if hosts is None:                                         # a whole network: find the devices first
        from src.net_scan import scan

        progress(f"Finding the devices on {label}…")
        found = scan(label, config, should_continue=should_continue,
                     progress=lambda done, total: progress(f"Finding the devices on {label}… {done}/{total}"))
        stopped, timed_out = found.stopped, found.timed_out
        reports = [HostReport(ip=d.ip, name=d.name, vendor=d.vendor) for d in found.devices]
    else:
        reports = [HostReport(ip=ip) for ip in hosts]

    mine = own_ips()
    outside = [r for r in reports if r.ip not in mine]
    pairs = [(r, port) for r in outside for port in PORTS]
    if pairs and not stopped:
        progress(f"Checking {len(PORTS)} ports on {len(outside)} host{'s' if len(outside) != 1 else ''}…")
        pool = ThreadPoolExecutor(max_workers=min(WORKERS, len(pairs)))
        try:
            futures = {pool.submit(_port_open, r.ip, port): (r, port) for r, port in pairs}
            open_pairs = []
            for fut in as_completed(futures):
                if not should_continue():
                    stopped = True
                    break
                if time.monotonic() - started > TIME_LIMIT_SECONDS:
                    timed_out = True
                    break
                if fut.result():
                    open_pairs.append(futures[fut])
            if not stopped and not timed_out and open_pairs:
                progress(f"Reading {len(open_pairs)} open services…")
                for report, service in pool.map(lambda p: (p[0], check_service(p[0].ip, p[1])), open_pairs):
                    report.services.append(service)
        finally:
            pool.shutdown(wait=False, cancel_futures=True)

    for r in reports:
        r.services.sort(key=lambda s: s.port)
        if r.ip in mine:                                       # this PC inside a network check: inside view
            r.local = local_view()
            r.local.api_port = config.api_port
            r.name = r.name or socket.gethostname()
    findings = [f for r in reports for f in (local_findings(r.local, r.ip) if r.local else host_findings(r))]
    reports.sort(key=_host_key)
    return SecurityResult(label, reports, sorted(findings, key=_order), round(time.monotonic() - started, 1),
                          stopped, timed_out)


# ── presentation ─────────────────────────────────────────────────────────
def _cell(text: str) -> str:
    return (text or "—").replace("|", "/").replace("\n", " ")


def _table(header: list[str], rows: list[list[str]]) -> str:
    lines = ["| " + " | ".join(header) + " |", "|" + "---|" * len(header)]
    return "\n".join(lines + ["| " + " | ".join(_cell(c) for c in row) + " |" for row in rows])


def finding_rows(result: SecurityResult) -> list[list[str]]:
    return [[RISK_LABEL[f.risk], f.host, str(f.port or "—"), f.title, f.why, f.fix] for f in result.findings]


def grouped_rows(result: SecurityResult) -> list[list[str]]:
    """Network checks: one row per kind of finding with all its hosts ("RTSP open — 8 hosts: …")."""
    groups: dict[tuple, list[str]] = {}
    for f in result.findings:
        title = re.sub(r" \([^)]*\)$", "", f.title)            # "SMB listening on the network (System)" → one kind
        groups.setdefault((f.risk, title, f.why, f.fix), []).append(f.host)
    rows = []
    for (risk, title, why, fix), hosts in groups.items():
        unique = sorted(dict.fromkeys(hosts), key=lambda h: _host_key(HostReport(ip=h)))
        listed = ", ".join(unique[:12]) + (f" … (+{len(unique) - 12})" if len(unique) > 12 else "")
        rows.append([RISK_LABEL[risk], title, f"{len(unique)} host{'s' if len(unique) != 1 else ''}: {listed}",
                     why, fix])
    return rows


def service_rows(result: SecurityResult) -> list[list[str]]:
    return [[r.ip, r.name or "—", str(s.port), s.name, s.product or "—", s.tls or "—"]
            for r in result.hosts if not r.local for s in r.services]


# names for common Windows ports that need no finding (display only)
KNOWN_LOCAL = {2179: "Hyper-V VM connection", 5040: "Windows Connected Devices", 5357: "Windows network discovery",
               7680: "Windows Update delivery"}


def _known_as(port: int, config_port: int | None = None) -> str:
    if port in PORTS:
        return PORTS[port].name
    if port == config_port:
        return "MulteAgent API"
    if port in KNOWN_LOCAL:
        return KNOWN_LOCAL[port]
    if 49152 <= port <= 65535:
        return "Windows dynamic port (RPC)"
    return "—"


def listener_rows(listeners: list[Listener], api_port: int | None = None) -> list[list[str]]:
    return [[lst.address, str(lst.port), lst.program, "yes" if lst.network else "no (this PC only)",
             _known_as(lst.port, api_port)] for lst in listeners]


def summary(result: SecurityResult) -> str:
    counts = {risk: sum(1 for f in result.findings if f.risk == risk) for risk in (HIGH, MEDIUM, LOW)}
    hosts = len(result.hosts)
    text = (f"Security check of **{result.target}** — {hosts} host{'s' if hosts != 1 else ''} in "
            f"{result.seconds:g} s: **🔴 {counts[HIGH]} high · 🟠 {counts[MEDIUM]} medium · 🟡 {counts[LOW]} low**.")
    if result.stopped:
        text += " _Stopped early: the check is incomplete._"
    if result.timed_out:
        text += f" _Time limit ({TIME_LIMIT_SECONDS} s) reached: the check is incomplete._"
    return text


NOTES = ("Read-only check: nothing was changed. From outside, a closed port can mean 'nothing runs' or 'a firewall "
         "blocks it'; versions come from what services announce and may be hidden. This PC is checked from the inside "
         "(scanning its own address would bypass Windows Firewall). The fixes are advice — run them yourself.")


def _onoff(value: bool | None) -> str:
    return "unknown" if value is None else ("**on**" if value else "off")


def markdown(result: SecurityResult) -> str:
    parts = [summary(result)]
    if not result.findings:
        parts.append("No risky ports or settings found by the rules.")
    elif len(result.hosts) > 1:
        parts.append("### Findings (grouped — every host and port is in the Excel file)\n" + _table(
            ["Risk", "Finding", "Hosts", "Why it matters", "What to do"], grouped_rows(result)))
    else:
        parts.append("### Findings\n" + _table(["Risk", "Host", "Port", "Finding", "Why it matters", "What to do"],
                                                finding_rows(result)))
    services = service_rows(result)
    if services and len(services) <= MAX_SERVICE_ROWS:
        parts.append("### Open services\n" + _table(["Host", "Name", "Port", "Service", "Version / product",
                                                     "Certificate"], services))
    elif services:
        parts.append(f"**{len(services)} open services** — the full list (host, port, software version, certificate) "
                     "is in the Excel file.")
    for r in result.hosts:
        if not r.local:
            continue
        view = r.local
        fw = ", ".join(f"{name} {'on' if on else '**OFF**'}" for name, on in view.firewall.items()) or "not read"
        lines = [f"### This PC ({r.name or r.ip})",
                 f"Windows Firewall: {fw} · Remote Desktop: {_onoff(view.rdp_on)} · SMBv1: {_onoff(view.smb1_on)}"]
        exposed = [lst for lst in view.listeners if lst.network]
        if exposed:
            lines.append(_table(["Address", "Port", "Program", "Reachable from network", "Known as"],
                                listener_rows(exposed, view.api_port)))
        lines.append(f"{len(view.listeners) - len(exposed)} more ports listen for this PC only (127.0.0.1) — "
                     "not reachable from other devices.")
        lines += [f"_{note}_" for note in view.notes]
        parts.append("\n\n".join(lines))
    parts.append(f"_{NOTES}_")
    return "\n\n".join(parts)


def xlsx_bytes(result: SecurityResult) -> bytes:
    from io import BytesIO

    from openpyxl import Workbook
    from openpyxl.styles import Font

    wb = Workbook()
    sheets = [("Findings", ["Risk", "Host", "Port", "Finding", "Why it matters", "What to do"], finding_rows(result),
               (10, 16, 7, 40, 50, 60)),
              ("Open services", ["Host", "Name", "Port", "Service", "Version / product", "Certificate"],
               service_rows(result), (16, 24, 7, 20, 40, 40))]
    for r in result.hosts:
        if r.local:
            sheets.append(("This PC", ["Address", "Port", "Program", "Reachable from network", "Known as"],
                           listener_rows(r.local.listeners, r.local.api_port), (18, 8, 28, 22, 20)))
    for i, (title, header, body, widths) in enumerate(sheets):
        ws = wb.active if i == 0 else wb.create_sheet()
        ws.title = title
        ws.append(header)
        for row in body:
            ws.append([re.sub(r"^[🔴🟠🟡⚪] ", "", c) for c in row])
        for cell in ws[1]:
            cell.font = Font(bold=True)
        for col, width in zip("ABCDEF", widths):
            ws.column_dimensions[col].width = width
        ws.freeze_panes = "A2"
    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()


# ── as a chat job ────────────────────────────────────────────────────────
def run_security_check(
    query: str,
    *,
    job_id: str,
    config: Config,
    on_progress: Callable[[dict], None] | None = None,
    target: str | None = None,
    **_: object,
) -> tuple[str, list[str], float]:
    """'check the security of 192.168.1.10' / 'is this PC secure?' → findings table + .xlsx. No model."""
    from src import live
    from src.artifacts import artifact_event, write_bytes

    emit = on_progress or (lambda event: None)
    agent = {"agent_id": "assistant", "name": "Assistant", "role": "Security check", "stage": "chat",
             "accent": "#c45c5c", "job_id": job_id}
    started = time.monotonic()
    errors: list[str] = []
    emit({"type": "session_start", "mode": "chat", "query": query})
    emit({"type": "agent_start", **agent, "model": "security_check"})

    def status(message: str) -> None:
        emit({"type": "status", "agent_id": "assistant", "message": message})

    try:
        result = check(target, query, config, should_continue=lambda: not live.is_stopped(job_id), progress=status)
        output = markdown(result)
        name = re.sub(r"[^0-9A-Za-z.]+", "_", result.target).strip("_")
        meta = write_bytes(config, job_id, f"security-check-{name}.xlsx", xlsx_bytes(result), kind="xlsx")
        emit(artifact_event(meta, job_id=job_id))
        if result.stopped:
            agent["role"] = "Stopped"
    except ScanError as exc:
        output = f"The security check was not run: {exc}."
        errors.append(str(exc))
    except Exception as exc:  # noqa: BLE001 — report, never crash the worker
        logger.exception("security check failed")
        output = f"The security check failed: {exc}"
        errors.append(str(exc))
    elapsed = time.monotonic() - started
    emit({"type": "agent_done", **agent, "elapsed_seconds": elapsed, "output": output})
    emit({"type": "session_done", "total_seconds": elapsed, "errors": errors, "mode": "chat"})
    return output, errors, elapsed
